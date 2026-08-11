#!/usr/bin/env python3
"""Add user-facing work-goal categories to the scenario database."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "发现页内容数据库.xlsx"

WORK_GOALS = {
    "信息提取与补全": {"S001", "S005", "S026", "S028"},
    "内容撰写与总结": {"S022", "S027", "S038", "S040", "S041"},
    "计划拆解与协同": {"S002", "S004", "S013", "S014", "S015", "S016", "S023", "S025"},
    "分类分派与流转": {"S003", "S006", "S007", "S009"},
    "质量检查与审核": {
        "S008", "S010", "S011", "S012", "S019", "S020",
        "S024", "S029", "S035", "S037", "S039",
    },
    "分析评估与预警": {"S021", "S030", "S031", "S032", "S033", "S034", "S036"},
}


def header_positions(sheet):
    return {cell.value: cell.column for cell in sheet[1]}


def main():
    workbook = load_workbook(SOURCE)
    scenarios = workbook["场景表"]
    positions = header_positions(scenarios)

    if "工作目标" not in positions:
        insert_at = positions["场景名称"] + 1
        scenarios.insert_cols(insert_at)
        scenarios.cell(1, insert_at, "工作目标")
        positions = header_positions(scenarios)

    goal_by_id = {
        scenario_id: goal
        for goal, scenario_ids in WORK_GOALS.items()
        for scenario_id in scenario_ids
    }
    scenario_ids = set()
    for row in range(2, scenarios.max_row + 1):
        scenario_id = scenarios.cell(row, positions["场景ID"]).value
        scenario_ids.add(scenario_id)
        scenarios.cell(row, positions["工作目标"], goal_by_id.get(scenario_id, ""))

    expected_method_ids = {
        scenario_id
        for scenario_id in scenario_ids
        if scenario_id not in {"S017", "S018"}
    }
    assert set(goal_by_id) == expected_method_ids

    goal_values = ",".join(WORK_GOALS)
    validation = DataValidation(
        type="list",
        formula1=f'"{goal_values}"',
        allow_blank=True,
    )
    scenarios.add_data_validation(validation)
    validation.add(
        f"{scenarios.cell(2, positions['工作目标']).coordinate}:"
        f"{scenarios.cell(300, positions['工作目标']).coordinate}"
    )

    fields = workbook["字段说明"]
    field_rows = {
        (fields.cell(row, 1).value, fields.cell(row, 2).value)
        for row in range(2, fields.max_row + 1)
    }
    if ("场景表", "工作目标") not in field_rows:
        fields.append([
            "场景表",
            "工作目标",
            "好方法必填",
            "按用户要完成的工作分类，不按 AI 应用、插件或轻应用等产品类型分类。",
        ])

    enums = workbook["枚举与规范"]
    enum_names = {
        enums.cell(row, 1).value
        for row in range(2, enums.max_row + 1)
    }
    if "工作目标" not in enum_names:
        enums.append([
            "工作目标",
            "｜".join(WORK_GOALS),
            "用于第二套方案“好方法”的用户侧分类。",
        ])

    workbook.save(SOURCE)
    print({
        "classified": len(goal_by_id),
        "categories": {goal: len(ids) for goal, ids in WORK_GOALS.items()},
        "source": str(SOURCE),
    })


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Add user-requested AI workflow scenarios to the content workbook."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "发现页内容数据库.xlsx"

SCENARIOS = [
    {
        "title": "撰写产品 PRD",
        "description": "需求进入方案阶段后，AI 汇总背景、目标、范围、用户故事和验收标准，生成结构化 PRD 初稿。",
        "problem": "产品经理需要从访谈、会议纪要和零散需求中反复整理 PRD，文档结构与信息完整度依赖个人经验。",
        "detail": "需求进入方案阶段时，读取当前需求及关联的访谈记录、会议纪要和参考资料，按团队模板生成包含背景、目标、范围、用户故事、方案说明、验收标准和风险项的 PRD 初稿，并标记待确认信息。",
        "value": "缩短文档起草时间｜统一 PRD 结构｜减少关键信息遗漏",
        "industry": "互联网｜消费电子｜零售",
        "role": "产品经理",
        "tags": "需求文档｜结构化写作｜PRD",
        "product_id": "A011",
        "product_name": "AI 生成云文档",
        "steps": "选择需求节点｜关联背景材料｜配置 PRD 模板｜生成并人工确认",
        "output": "结构化 PRD 初稿",
        "prompt": "你是产品需求文档助手。请结合当前需求、访谈记录、会议纪要和关联资料，按照“背景与问题、目标与指标、用户与场景、功能范围、用户故事、方案说明、验收标准、风险与待确认项”的结构生成 PRD 初稿。仅使用已有材料，不明确的信息标记为“待确认”，不要虚构数据或结论。",
    },
    {
        "title": "走查产品需求",
        "description": "需求提交评审前，AI 按 PM 走查清单检查目标、场景、范围、逻辑和验收标准，输出修改建议。",
        "problem": "需求走查依赖产品经理个人经验，范围遗漏、逻辑断点和验收标准不清常在正式评审时才暴露。",
        "detail": "需求提交评审前，读取 PRD、原型和关联背景，按团队 PM 走查清单检查目标是否明确、用户场景是否闭环、范围边界是否清晰、异常流程是否覆盖、验收标准是否可验证，并将问题和修改建议回写到工作项。",
        "value": "前置发现需求问题｜统一走查标准｜提升正式评审效率",
        "industry": "互联网｜消费电子｜游戏",
        "role": "产品经理｜项目经理",
        "tags": "需求走查｜逻辑检查｜评审准备",
        "product_id": "A001",
        "product_name": "自定义指令",
        "steps": "配置走查清单｜关联 PRD 与原型｜执行需求检查｜确认修改建议",
        "output": "需求走查问题清单",
        "prompt": "你是资深产品经理，请在需求提交正式评审前完成 PM 走查。检查：1. 背景、目标和成功指标是否明确；2. 用户角色与核心场景是否完整；3. 功能范围和非目标是否清晰；4. 主流程、异常流程和权限逻辑是否闭环；5. 验收标准是否具体可验证。按“问题位置、问题描述、影响、修改建议、严重程度”输出，不替作者补写未经确认的业务事实。",
    },
    {
        "title": "总结项目成果",
        "description": "项目或阶段结束后，AI 汇总目标达成、关键交付、风险处理和经验教训，生成项目总结初稿。",
        "problem": "项目总结依赖人工跨任务、文档和会议记录取数，容易只罗列过程，缺少目标对照和可复用经验。",
        "detail": "项目进入结项或阶段验收节点后，读取目标、里程碑、交付物、风险、变更和关键决策，生成包含目标达成情况、关键成果、问题与处理、经验教训和后续行动的项目总结，并保留信息来源。",
        "value": "降低总结成本｜沉淀项目经验｜形成后续行动闭环",
        "industry": "互联网｜制造｜教育｜能源",
        "role": "项目经理｜高管",
        "tags": "项目总结｜成果沉淀｜经验复用",
        "product_id": "A005",
        "product_name": "内容总结",
        "steps": "选择结项节点｜汇总项目数据｜配置总结结构｜生成并确认",
        "output": "项目总结初稿",
        "prompt": "你是项目总结助手。请结合项目目标、里程碑、交付物、风险记录、变更记录和关键决策，生成项目总结。结构包括：目标达成情况、关键成果、未完成事项、主要问题与处理、有效做法、经验教训、后续行动。每项结论注明依据；无法从材料确认的内容标记“待补充”，不要夸大成果。",
    },
    {
        "title": "撰写产品手册",
        "description": "版本验收后，AI 汇总功能说明、操作步骤和注意事项，生成面向用户的产品手册初稿。",
        "problem": "产品手册需要人工汇总需求、原型和操作说明，更新经常滞后于版本，用户难以及时理解新功能。",
        "detail": "版本进入验收或发布节点后，读取已交付需求、功能说明、操作步骤、权限限制和常见问题，按产品手册目录生成面向用户的云文档，并标记缺失截图和待确认内容。",
        "value": "缩短手册编写时间｜保持文档与版本一致｜降低用户学习成本",
        "industry": "互联网｜消费电子｜教育",
        "role": "产品经理｜运营",
        "tags": "产品手册｜用户文档｜知识沉淀",
        "product_id": "A011",
        "product_name": "AI 生成云文档",
        "steps": "选择发布节点｜关联版本材料｜配置手册目录｜生成并人工校对",
        "output": "产品手册初稿",
        "prompt": "你是产品文档助手。请结合当前版本的需求说明、功能描述、操作步骤、权限规则和常见问题，按照“产品概述、适用对象、功能说明、操作指引、权限与限制、常见问题、版本信息”的结构生成产品手册初稿。使用面向最终用户的语言；缺失截图或无法确认的信息标记为“待补充”。",
    },
]

TEST_UPDATE = {
    "卡片短描述": "需求进入测试阶段后，AI 根据需求描述和验收标准生成覆盖主流程、边界条件和异常场景的测试用例。",
    "问题描述": "测试人员需要逐条阅读需求并设计用例，重复劳动多，边界条件和异常流程容易遗漏。",
    "场景说明": "需求进入测试阶段时，读取 PRD、验收标准和关联原型，按团队用例模板生成前置条件、操作步骤、预期结果和优先级，并标记需要人工补充的特殊场景。",
    "预期价值": "缩短用例设计时间｜提升场景覆盖率｜统一测试用例格式",
    "场景标签": "测试用例｜边界分析｜质量保障",
}

TEST_RELATION_UPDATE = {
    "使用方法概述": "在需求进入测试阶段时生成结构化测试用例，并写入测试管理流程。",
    "操作步骤": "选择测试触发节点｜关联 PRD 与原型｜配置用例模板｜生成并评审",
    "参考提示词": "你是测试用例设计助手。请结合当前需求、验收标准和关联原型，生成结构化测试用例，覆盖主流程、分支流程、边界条件、异常输入、权限差异和状态转换。每条用例包含用例名称、前置条件、操作步骤、预期结果和优先级；无法确认的规则标记为“待确认”，不要虚构系统行为。",
    "预期产出": "结构化测试用例",
}


def headers(sheet):
    return [cell.value for cell in sheet[1]]


def next_id(prefix, values):
    numbers = [
        int(value[len(prefix):])
        for value in values
        if value and value.startswith(prefix) and value[len(prefix):].isdigit()
    ]
    return f"{prefix}{max(numbers, default=0) + 1:03d}"


def append_dict(sheet, row):
    names = headers(sheet)
    sheet.append([row.get(name, "") for name in names])


def update_row(sheet, id_column, id_value, values):
    names = headers(sheet)
    positions = {name: index for index, name in enumerate(names)}
    for row in sheet.iter_rows(min_row=2):
        if row[positions[id_column]].value == id_value:
            for name, value in values.items():
                row[positions[name]].value = value
            return
    raise RuntimeError(f"{id_column}={id_value} not found in {sheet.title}")


def main():
    workbook = load_workbook(WORKBOOK)
    scenarios_sheet = workbook["场景表"]
    relations_sheet = workbook["场景产品关系表"]
    intake_sheet = workbook["案例录入表"]

    scenario_rows = [
        dict(zip(headers(scenarios_sheet), row))
        for row in scenarios_sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]
    relation_rows = [
        dict(zip(headers(relations_sheet), row))
        for row in relations_sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]
    titles = {row["场景名称"] for row in scenario_rows}
    scenario_ids = {row["场景ID"] for row in scenario_rows}
    relation_ids = {row["关系ID"] for row in relation_rows}
    sort_order = max(row["推荐排序"] or 0 for row in scenario_rows)

    test_scene = next(row for row in scenario_rows if row["场景名称"] == "生成测试用例")
    update_row(scenarios_sheet, "场景ID", test_scene["场景ID"], TEST_UPDATE)
    test_relation = next(row for row in relation_rows if row["场景ID"] == test_scene["场景ID"])
    update_row(relations_sheet, "关系ID", test_relation["关系ID"], TEST_RELATION_UPDATE)

    created = []
    for item in SCENARIOS:
        if item["title"] in titles:
            continue
        scenario_id = next_id("S", scenario_ids)
        scenario_ids.add(scenario_id)
        relation_id = next_id("R", relation_ids)
        relation_ids.add(relation_id)
        sort_order += 1

        append_dict(scenarios_sheet, {
            "场景ID": scenario_id,
            "场景名称": item["title"],
            "卡片短描述": item["description"],
            "问题描述": item["problem"],
            "场景说明": item["detail"],
            "预期价值": item["value"],
            "适用行业": item["industry"],
            "适用角色": item["role"],
            "场景标签": item["tags"],
            "封面类型": "流程节点",
            "是否推荐": "是",
            "推荐排序": sort_order,
            "内容状态": "模拟待校准",
            "真实性备注": "用户指定新增场景；描述、提示词及映射为模拟内容，待产品能力确认。",
        })
        append_dict(relations_sheet, {
            "关系ID": relation_id,
            "场景ID": scenario_id,
            "产品ID": item["product_id"],
            "是否主要实现": "是",
            "使用方法概述": f"在对应流程节点中配置{item['product_name']}，自动完成{item['title']}。",
            "操作步骤": item["steps"],
            "交付类型": "参考提示词",
            "参考提示词": item["prompt"],
            "按钮文案": "复制参考提示词",
            "配置注意事项": "示例内容；上线前需校验字段、权限、触发条件和人工复核机制。",
            "预期产出": item["output"],
            "内容状态": "模拟待校准",
        })
        created.append((scenario_id, item))
        titles.add(item["title"])

    intake_headers = headers(intake_sheet)
    intake_positions = {name: index + 1 for index, name in enumerate(intake_headers)}
    existing_intake_titles = {
        row[intake_positions["场景名称"] - 1]
        for row in intake_sheet.iter_rows(min_row=2, values_only=True)
        if row[intake_positions["场景名称"] - 1]
    }
    intake_by_scene = {scenario_id: item for scenario_id, item in created}
    intake_by_scene[test_scene["场景ID"]] = {
        "title": "生成测试用例",
        "description": TEST_UPDATE["场景说明"],
        "product_id": "A009",
        "product_name": "智能测试用例生成",
        "industry": test_scene["适用行业"],
        "role": test_scene["适用角色"],
    }
    empty_rows = [
        row_number
        for row_number in range(2, intake_sheet.max_row + 1)
        if not intake_sheet.cell(row_number, intake_positions["场景名称"]).value
    ]
    for scenario_id, item in intake_by_scene.items():
        if item["title"] in existing_intake_titles:
            continue
        row_number = empty_rows.pop(0)
        values = {
            "关联场景ID（可选）": scenario_id,
            "场景名称": item["title"],
            "原始场景描述": item["description"],
            "AI应用ID（可选）": item["product_id"],
            "AI应用名称": item["product_name"],
            "适用行业": item["industry"],
            "适用角色": item["role"],
            "真实性备注": "用户指定新增场景；文案由助手整理。",
        }
        for name, value in values.items():
            intake_sheet.cell(row_number, intake_positions[name], value)

    workbook.save(WORKBOOK)
    print({
        "created": [item["title"] for _, item in created],
        "updated": ["生成测试用例"],
        "workbook": str(WORKBOOK),
    })


if __name__ == "__main__":
    main()

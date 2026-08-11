#!/usr/bin/env python3
"""Apply the reviewed AI-node scenario selection without rebuilding the workbook."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "发现页内容数据库.xlsx"

SELECTED_AI_SCENARIOS = {
    "S001",
    "S002",
    "S003",
    "S004",
    "S005",
    "S010",
    "S012",
}

ASSISTANT_TEMPLATE_SCENARIOS = {
    "S006",
    "S007",
    "S008",
    "S009",
    "S011",
}

SCENARIO_UPDATES = {
    "S002": {
        "卡片短描述": "根据目标、范围和约束自动拆成可执行任务，补全验收标准、依赖和负责人建议。",
        "问题描述": "复杂需求依赖人工逐层拆解，任务粒度和验收口径不一致，关键依赖容易遗漏。",
        "场景说明": "需求进入规划节点后，读取目标、范围、交付时间和团队约束，生成任务清单、验收标准、依赖关系与负责人建议，由项目负责人确认后写入项目。",
        "预期价值": "缩短任务拆解时间｜统一任务粒度｜提前暴露依赖",
        "场景标签": "任务拆解｜项目规划｜AI节点",
    },
    "S012": {
        "卡片短描述": "汇总需求、功能说明和操作步骤，自动生成结构化产品手册草稿。",
        "问题描述": "产品手册依赖人工汇总多处材料，更新不及时，功能说明与实际版本容易脱节。",
        "场景说明": "版本进入发布或验收节点后，读取关联需求、功能说明、操作步骤和注意事项，按固定目录生成产品手册云文档，并标记缺失信息供人工补充。",
        "预期价值": "缩短手册编写时间｜保持版本信息一致｜沉淀可复用产品知识",
        "适用行业": "互联网｜消费电子｜游戏",
        "适用角色": "产品经理｜运营｜研发",
        "场景标签": "文档生成｜产品知识｜AI节点",
    },
}

RELATION_UPDATES = {
    "S001": {
        "产品ID": "A008",
        "使用方法概述": "在流程节点中提取关联材料并回填客户字段。",
        "操作步骤": "选择触发节点｜配置材料来源｜映射输出字段｜用样例验证",
        "预期产出": "结构化客户信息",
    },
    "S002": {
        "产品ID": "A010",
        "使用方法概述": "在需求进入规划阶段时生成任务清单、验收标准和依赖建议。",
        "操作步骤": "选择触发节点｜配置需求上下文｜定义任务输出格式｜审核后写入",
        "预期产出": "任务清单与依赖建议",
    },
    "S003": {
        "产品ID": "A003",
        "使用方法概述": "在客户反馈进入流程时自动分类、打标并触发后续路由。",
        "操作步骤": "定义分类标签｜配置判断上下文｜设置后续路由｜抽样复核",
        "预期产出": "反馈标签与路由结果",
    },
    "S004": {
        "产品ID": "A010",
        "使用方法概述": "连接项目上下文，生成任务树、依赖关系和初始排期。",
        "操作步骤": "输入项目目标｜补充资源约束｜生成计划草案｜审核后写入",
        "预期产出": "任务树与初始排期",
    },
    "S005": {
        "产品ID": "A001",
        "使用方法概述": "在提测或发布节点按团队规则检查交付材料完整性。",
        "操作步骤": "配置检查规则｜选择关联材料｜运行质量自检｜处理整改清单",
        "预期产出": "交付质量自检清单",
    },
    "S010": {
        "产品ID": "A004",
        "使用方法概述": "在合同进入评审或归档节点时提取关键条款并写入项目字段。",
        "操作步骤": "选择合同材料｜定义提取字段｜配置风险规则｜确认后写入",
        "预期产出": "合同关键条款与履约风险",
    },
    "S012": {
        "产品ID": "A011",
        "使用方法概述": "在版本发布或验收节点汇总关联材料并生成产品手册云文档。",
        "操作步骤": "选择触发节点｜配置材料范围｜定义手册目录｜生成并人工校对",
        "预期产出": "产品手册云文档草稿",
    },
}

REFERENCE_PROMPTS = {
    "S001": """你是客户信息整理助手。请阅读当前工作项及关联的会议纪要、邮件和附件，提取客户名称、所属行业、联系人、核心诉求、期望交付时间和待确认事项。仅使用材料中明确出现的信息；无法确认的字段填写“待人工确认”。请按字段名、提取结果、信息来源输出结构化结果。""",
    "S002": """你是项目任务拆解助手。请根据当前需求的目标、范围、交付时间和团队约束，将需求拆解为可执行任务。每项任务需包含任务名称、任务说明、验收标准、前置依赖和负责人角色建议。不要虚构未提供的技术方案；缺失条件请单独列为待确认项。""",
    "S003": """你是客户反馈分类助手。请分析当前反馈，将其归入预设的反馈类型和产品模块，并判断优先级。输出反馈类型、产品模块、优先级、判断理由、是否疑似重复及建议流转队列。信息不足时标记“待人工确认”，不要自行补充事实。""",
    "S004": """你是项目规划助手。请结合项目目标、范围、资源约束和交付日期，生成分层任务树、关键依赖、里程碑和初始排期建议。说明每项建议的依据，并将资源冲突、外部依赖和信息缺口列入风险清单。输出内容将由项目经理审核后写入项目。""",
    "S005": """你是交付质量检查助手。请按团队质量门禁规则检查当前工作项关联的需求说明、代码提交、测试用例、测试结果、验收记录和发布说明。输出已满足项、缺失项、风险等级和具体整改建议。不得将未找到的材料视为已完成。""",
    "S010": """你是合同信息提取助手。请阅读当前合同材料，提取合同主体、金额、期限、交付物、关键里程碑、验收条件、付款条件、违约责任和终止条款。标注每项内容的原文依据；对存在歧义或未出现的内容标记“待人工确认”。同时列出可能影响项目履约的风险点。""",
    "S012": """你是产品文档助手。请汇总当前版本关联的需求说明、功能描述、操作步骤和注意事项，按照“产品概述、适用对象、功能说明、操作指引、限制与注意事项、版本信息”的目录生成产品手册草稿。仅基于已有材料撰写，并将缺失内容标记为待补充。""",
}


def sheet_rows(ws):
    headers = [cell.value for cell in ws[1]]
    return headers, {
        row[0].value: (index, row)
        for index, row in enumerate(ws.iter_rows(min_row=2), start=2)
        if row[0].value
    }


def update_cells(headers, row, values):
    positions = {name: index for index, name in enumerate(headers)}
    for name, value in values.items():
        row[positions[name]].value = value


def main():
    workbook = load_workbook(SOURCE)

    scenario_sheet = workbook["场景表"]
    scenario_headers, scenario_rows = sheet_rows(scenario_sheet)
    for scenario_id, values in SCENARIO_UPDATES.items():
        update_cells(scenario_headers, scenario_rows[scenario_id][1], values)
    for scenario_id in SELECTED_AI_SCENARIOS:
        update_cells(
            scenario_headers,
            scenario_rows[scenario_id][1],
            {"是否推荐": "是"},
        )
    for scenario_id in ASSISTANT_TEMPLATE_SCENARIOS:
        update_cells(
            scenario_headers,
            scenario_rows[scenario_id][1],
            {"是否推荐": "否"},
        )

    relation_sheet = workbook["场景产品关系表"]
    relation_headers = [cell.value for cell in relation_sheet[1]]
    positions = {name: index for index, name in enumerate(relation_headers)}
    primary_relations = {}
    for row in relation_sheet.iter_rows(min_row=2):
        scenario_id = row[positions["场景ID"]].value
        is_primary = row[positions["是否主要实现"]].value
        if scenario_id in SELECTED_AI_SCENARIOS and is_primary == "是":
            primary_relations[scenario_id] = row

    missing = SELECTED_AI_SCENARIOS - primary_relations.keys()
    if missing:
        raise RuntimeError(f"Missing primary relations: {sorted(missing)}")

    for scenario_id, values in RELATION_UPDATES.items():
        update_cells(relation_headers, primary_relations[scenario_id], values)
        update_cells(
            relation_headers,
            primary_relations[scenario_id],
            {"参考提示词": REFERENCE_PROMPTS[scenario_id]},
        )

    workbook.save(SOURCE)
    print(
        {
            "selected_ai_scenarios": sorted(SELECTED_AI_SCENARIOS),
            "hidden_assistant_templates": sorted(ASSISTANT_TEMPLATE_SCENARIOS),
            "source": str(SOURCE),
        }
    )


if __name__ == "__main__":
    main()

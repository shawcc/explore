#!/usr/bin/env python3
"""Import selected workflow-friendly AI scenarios from the user-curated CSV."""

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CSV_SOURCE = ROOT / "Meego发现_Meego发现_表格.csv"
WORKBOOK = ROOT / "发现页内容数据库.xlsx"
SEP = "｜"

# Selected by workflow-node fit: explicit trigger, bounded input, actionable output.
SELECTED = {
    1: "补全客户信息",
    2: "拆解会议任务",
    3: "创建外部工单",
    4: "创建邮件任务",
    5: "提取合同条款",
    6: "分类客户反馈",
    7: "分派研发需求",
    8: "评定缺陷等级",
    9: "分流客服工单",
    10: "检查交付质量",
    11: "检查需求合规",
    12: "检查代码提交",
    13: "检查设计规范",
    16: "评估变更影响",
    17: "预警供应链风险",
    20: "起草发布说明",
    23: "拆解项目计划",
    28: "生成测试用例",
    30: "响应系统故障",
    82: "补全需求描述",
    83: "生成缺陷报告",
    84: "补全企业信息",
    89: "审查金融合规",
    90: "评估投资风险",
    93: "评估销售线索",
    101: "预警项目延期",
    102: "预警预算超支",
    104: "预警客户流失",
    109: "检查申报材料",
    110: "评估门店选址",
    111: "审核报名资格",
}

EXISTING_PRODUCT_ALIASES = {
    "AI 智能填单": "A008",
    "AI 智能会议分析": "A012",
    "AI 测试生成": "A009",
}

VALUE_RULES = [
    (("补全", "创建", "提取", "生成", "起草"), "减少人工处理｜统一输出结构｜降低信息遗漏"),
    (("分类", "分流", "分派", "评定"), "缩短流转时间｜提升判断一致性｜减少人工分派"),
    (("检查", "审查", "审核"), "前置发现问题｜统一审核标准｜降低合规风险"),
    (("评估",), "缩短评估时间｜提供决策依据｜暴露关键风险"),
    (("预警",), "提前识别风险｜缩短响应时间｜推动责任闭环"),
    (("响应",), "缩短发现到响应时间｜自动沉淀上下文｜加快问题闭环"),
    (("拆解",), "统一拆解口径｜减少遗漏｜加快进入执行"),
]


def read_csv():
    with CSV_SOURCE.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    by_number = {int(row["序号"]): row for row in rows}
    missing = set(SELECTED) - by_number.keys()
    if missing:
        raise RuntimeError(f"CSV is missing selected rows: {sorted(missing)}")
    return [by_number[number] for number in SELECTED]


def headers(sheet):
    return [cell.value for cell in sheet[1]]


def dict_rows(sheet):
    names = headers(sheet)
    return [
        dict(zip(names, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]


def replace_rows(sheet, names, rows):
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append([row.get(name, "") for name in names])


def tags(row):
    return SEP.join(
        part.strip()
        for part in re.split(r"\s*[·｜]\s*", row["meta"])
        if part.strip()
    )


def expected_value(title):
    for keywords, value in VALUE_RULES:
        if any(keyword in title for keyword in keywords):
            return value
    return "减少人工处理｜提升执行一致性｜沉淀结构化结果"


def problem_text(title, row):
    action = title[:2]
    return (
        f"当前“{title}”依赖人工读取材料、判断规则并执行后续操作，"
        f"处理效率和结果一致性容易受个人经验影响。"
    )


def scenario_text(row):
    description = row["描述(desc)"].strip()
    return re.sub(
        r"[，,](让|帮助|确保).*$",
        "。",
        description,
    ).replace("AI 自动", "AI 节点自动")


def card_text(row):
    text = scenario_text(row)
    text = re.sub(r"^(需求评审后|需求评审会结束后|上传合同文件后|客服收到用户反馈后)，?", "", text)
    return text[:76].rstrip("，,。") + "。"


def prompt_text(title, row):
    return (
        f"你是项目流程中的“{title}”助手。请结合当前工作项、关联字段和附件完成以下任务："
        f"1. 按场景规则处理输入信息；2. 输出可直接写回项目的结构化结果；"
        f"3. 标注判断依据和待人工确认项；4. 不补充材料中不存在的事实。"
        f"场景要求：{scenario_text(row)}"
    )


def next_id(prefix, used):
    number = max(
        [int(value[len(prefix):]) for value in used if value.startswith(prefix) and value[len(prefix):].isdigit()]
        or [0]
    ) + 1
    while f"{prefix}{number:03d}" in used:
        number += 1
    value = f"{prefix}{number:03d}"
    used.add(value)
    return value


def main():
    selected_rows = read_csv()
    workbook = load_workbook(WORKBOOK)

    scenario_sheet = workbook["场景表"]
    product_sheet = workbook["产品表"]
    relation_sheet = workbook["场景产品关系表"]
    intake_sheet = workbook["案例录入表"]

    scenario_headers = headers(scenario_sheet)
    product_headers = headers(product_sheet)
    relation_headers = headers(relation_sheet)

    scenarios = dict_rows(scenario_sheet)
    products = dict_rows(product_sheet)
    relations = dict_rows(relation_sheet)

    # Preserve the six non-AI examples. Rebuild the AI scenario set deterministically.
    non_ai_scenarios = [row for row in scenarios if row["场景ID"] in {f"S{i:03d}" for i in range(13, 19)}]
    non_ai_relations = [row for row in relations if row["场景ID"] in {f"S{i:03d}" for i in range(13, 19)}]

    product_by_name = {row["产品名称"]: row for row in products}
    product_by_id = {row["产品ID"]: row for row in products}
    used_product_ids = set(product_by_id)

    imported_scenarios = []
    imported_relations = []
    intake_records = []
    used_scenario_ids = {row["场景ID"] for row in non_ai_scenarios}
    used_relation_ids = {row["关系ID"] for row in non_ai_relations}

    for index, csv_row in enumerate(selected_rows, start=1):
        source_number = int(csv_row["序号"])
        title = SELECTED[source_number]
        scenario_id = f"S{index:03d}" if index <= 12 else next_id("S", used_scenario_ids)
        used_scenario_ids.add(scenario_id)

        feature_name = csv_row["功能名(featureName)"].strip() or "AI 助手"
        alias_id = EXISTING_PRODUCT_ALIASES.get(feature_name)
        product = product_by_id.get(alias_id) if alias_id else product_by_name.get(feature_name)
        if not product:
            product_id = next_id("A", used_product_ids)
            product = {name: "" for name in product_headers}
            product.update({
                "产品ID": product_id,
                "产品名称": feature_name,
                "产品类型": "AI应用",
                "产品子类型": "AI节点",
                "卡片短描述": f"在项目流程中执行“{title}”，并将结构化结果写回工作项。",
                "详情介绍": f"{feature_name}用于在明确的流程触发点执行{title}，输入来自当前工作项及关联材料，输出可用于字段回填、流程判断或后续任务。",
                "核心能力": tags(csv_row),
                "适用行业": csv_row["行业(industry)"],
                "适用角色": csv_row["角色(role)"],
                "封面图片": csv_row["图片URL(image)"],
                "提供方": "飞书项目（模拟）",
                "目录是否展示": "是",
                "目录排序": 100 + index,
                "内容状态": "模拟待校准",
                "真实性备注": f"名称和案例来自用户筛选 CSV 序号 {source_number}；产品能力与可用性待产品确认。",
            })
            products.append(product)
            product_by_name[feature_name] = product
            product_by_id[product_id] = product
        product_id = product["产品ID"]

        scenario = {name: "" for name in scenario_headers}
        scenario.update({
            "场景ID": scenario_id,
            "场景名称": title,
            "卡片短描述": card_text(csv_row),
            "问题描述": problem_text(title, csv_row),
            "场景说明": scenario_text(csv_row),
            "预期价值": expected_value(title),
            "适用行业": csv_row["行业(industry)"],
            "适用角色": csv_row["角色(role)"],
            "场景标签": tags(csv_row),
            "封面类型": "流程节点",
            "封面图片": csv_row["图片URL(image)"],
            "是否推荐": "是",
            "推荐排序": index,
            "内容状态": "模拟待校准",
            "真实性备注": f"来源：用户筛选 CSV 序号 {source_number}；场景名称和描述已按流程节点口径整理。",
        })
        imported_scenarios.append(scenario)

        relation_id = next_id("R", used_relation_ids)
        relation = {name: "" for name in relation_headers}
        relation.update({
            "关系ID": relation_id,
            "场景ID": scenario_id,
            "产品ID": product_id,
            "是否主要实现": "是",
            "使用方法概述": f"在对应流程节点中配置{feature_name}，自动完成{title}。",
            "操作步骤": "选择触发节点｜配置输入材料｜定义输出字段｜用样例验证",
            "交付类型": "参考提示词",
            "参考提示词": prompt_text(title, csv_row),
            "按钮文案": "复制参考提示词",
            "配置注意事项": "示例内容；上线前需校验字段、权限、触发条件和人工复核机制。",
            "预期产出": title,
            "内容状态": "模拟待校准",
        })
        imported_relations.append(relation)

        intake_records.append({
            "案例ID": f"CASE-{index:04d}",
            "关联场景ID（可选）": scenario_id,
            "场景名称": title,
            "原始场景描述": csv_row["描述(desc)"],
            "AI应用ID（可选）": product_id,
            "AI应用名称": feature_name,
            "适用行业": csv_row["行业(industry)"],
            "适用角色": csv_row["角色(role)"],
            "来源URL": "",
            "真实性备注": f"用户筛选 CSV 序号 {source_number}；图片：{csv_row['图片URL(image)']}",
        })

    replace_rows(scenario_sheet, scenario_headers, imported_scenarios + non_ai_scenarios)
    replace_rows(product_sheet, product_headers, products)
    replace_rows(relation_sheet, relation_headers, imported_relations + non_ai_relations)

    intake_headers = headers(intake_sheet)
    for row_number in range(2, intake_sheet.max_row + 1):
        for column in range(2, 11):
            intake_sheet.cell(row_number, column, None)
    for row_number, record in enumerate(intake_records, start=2):
        for column, name in enumerate(intake_headers, start=1):
            if name in record:
                intake_sheet.cell(row_number, column, record[name])

    workbook.save(WORKBOOK)
    print({
        "selected_ai_nodes": len(imported_scenarios),
        "products_total": len(products),
        "relations_total": len(imported_relations + non_ai_relations),
        "workbook": str(WORKBOOK),
    })


if __name__ == "__main__":
    main()

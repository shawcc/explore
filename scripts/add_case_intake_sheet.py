#!/usr/bin/env python3
"""Create or migrate the scene-first case intake sheet."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "发现页内容数据库.xlsx"
SHEET_NAME = "案例录入表"
DICTIONARY_SHEET = "录入字典"
MAX_CASES = 200

HEADERS = [
    "案例ID",
    "关联场景ID（可选）",
    "场景名称",
    "原始场景描述",
    "AI应用ID（可选）",
    "AI应用名称",
    "适用行业",
    "适用角色",
    "来源URL",
    "真实性备注",
    "匹配状态",
]

WIDTHS = [15, 19, 26, 54, 18, 25, 28, 28, 44, 38, 20]

COMMENTS = {
    "案例ID": "稳定主键，已预生成。请勿修改或重复使用。",
    "关联场景ID（可选）": "更新已有场景时选择；新增场景时留空。",
    "场景名称": "必填。填写用户能理解的业务场景，不填产品功能名。",
    "原始场景描述": "填写真实案例、问题、输入材料、执行过程和产出，不必润色。",
    "AI应用ID（可选）": "已有 AI 应用时从下拉列表选择；新 AI 应用留空。",
    "AI应用名称": "必填。已有应用需与所选 ID 对应；新应用直接填写名称。",
    "适用行业": "多选使用全角竖线｜分隔，例如：互联网｜金融。",
    "适用角色": "好方法可填写；新场景可留空。多选使用全角竖线｜分隔。",
    "来源URL": "填写可追溯的产品页、文档或案例链接。",
    "真实性备注": "说明哪些是公开事实、内部案例、模拟内容或待确认信息。",
    "匹配状态": "自动计算。同步程序只按 ID 更新，不按名称猜测。",
}


def rows_as_dicts(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [
        dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]


def rebuild_dictionary(workbook):
    if DICTIONARY_SHEET in workbook.sheetnames:
        del workbook[DICTIONARY_SHEET]
    sheet = workbook.create_sheet(DICTIONARY_SHEET)
    sheet.append(["场景ID", "场景名称", "AI应用ID", "AI应用名称"])

    scenario_sheet = workbook["场景表"]
    scenario_headers = [cell.value for cell in scenario_sheet[1]]
    scenarios = [
        dict(zip(scenario_headers, row))
        for row in scenario_sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]

    product_sheet = workbook["产品表"]
    product_headers = [cell.value for cell in product_sheet[1]]
    ai_products = [
        dict(zip(product_headers, row))
        for row in product_sheet.iter_rows(min_row=2, values_only=True)
        if row[0] and row[2] == "AI应用"
    ]

    for index in range(max(len(scenarios), len(ai_products))):
        scenario = scenarios[index] if index < len(scenarios) else {}
        product = ai_products[index] if index < len(ai_products) else {}
        sheet.append([
            scenario.get("场景ID", ""),
            scenario.get("场景名称", ""),
            product.get("产品ID", ""),
            product.get("产品名称", ""),
        ])

    sheet.sheet_state = "hidden"


def create_intake_sheet(workbook, existing_rows):
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME, 1)
    sheet.append(HEADERS)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{MAX_CASES + 1}"
    sheet.sheet_view.showGridLines = False

    fill = PatternFill("solid", fgColor="EDE8FF")
    for index, cell in enumerate(sheet[1], start=1):
        cell.fill = fill
        cell.font = Font(bold=True, color="4F3692")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.comment = Comment(COMMENTS[cell.value], "TRAE")
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[index - 1]

    for row_number in range(2, MAX_CASES + 2):
        sheet.cell(row_number, 1, f"CASE-{row_number - 1:04d}")
        ai_status = (
            f'IF(E{row_number}="",IF(F{row_number}="","待填AI应用","待建产品"),'
            f'IF(COUNTIF(\'{DICTIONARY_SHEET}\'!$C$2:$C$200,E{row_number})=0,"AI应用ID无效",'
            f'IF(F{row_number}="","补充AI应用名称",'
            f'IF(F{row_number}<>VLOOKUP(E{row_number},\'{DICTIONARY_SHEET}\'!$C$2:$D$200,2,FALSE),'
            f'"AI应用名称与ID不一致","可同步"))))'
        )
        sheet.cell(
            row_number,
            11,
            (
                f'=IF(COUNTA(C{row_number}:D{row_number},F{row_number})=0,"",'
                f'IF(B{row_number}="",{ai_status},'
                f'IF(COUNTIF(\'{DICTIONARY_SHEET}\'!$A$2:$A$200,B{row_number})=0,"场景ID无效",'
                f'IF(C{row_number}<>VLOOKUP(B{row_number},\'{DICTIONARY_SHEET}\'!$A$2:$B$200,2,FALSE),'
                f'"场景名称与ID不一致",{ai_status}))))'
            ),
        )
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    scenario_validation = DataValidation(
        type="list",
        formula1=f"'{DICTIONARY_SHEET}'!$A$2:$A$200",
        allow_blank=True,
    )
    ai_validation = DataValidation(
        type="list",
        formula1=f"'{DICTIONARY_SHEET}'!$C$2:$C$200",
        allow_blank=True,
    )
    sheet.add_data_validation(scenario_validation)
    sheet.add_data_validation(ai_validation)
    scenario_validation.add(f"B2:B{MAX_CASES + 1}")
    ai_validation.add(f"E2:E{MAX_CASES + 1}")

    for index, old in enumerate(existing_rows, start=2):
        if index > MAX_CASES + 1:
            break
        values = {
            2: old.get("关联场景ID（可选）", ""),
            3: old.get("场景名称", ""),
            4: old.get("原始场景描述", ""),
            5: old.get("AI应用ID（可选）", ""),
            6: old.get("AI应用名称", ""),
            7: old.get("适用行业", ""),
            8: old.get("适用角色", ""),
            9: old.get("来源URL", ""),
            10: old.get("真实性备注", ""),
        }
        for column, value in values.items():
            sheet.cell(index, column, value)


def main():
    workbook = load_workbook(SOURCE)
    existing_rows = (
        rows_as_dicts(workbook[SHEET_NAME])
        if SHEET_NAME in workbook.sheetnames
        else []
    )
    rebuild_dictionary(workbook)
    create_intake_sheet(workbook, existing_rows)
    workbook.save(SOURCE)
    print({
        "sheet": SHEET_NAME,
        "columns": HEADERS,
        "preserved_rows": len(existing_rows),
        "source": str(SOURCE),
    })


if __name__ == "__main__":
    main()

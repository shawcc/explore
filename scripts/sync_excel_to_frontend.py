#!/usr/bin/env python3
"""Compile the editable Excel content database into frontend JavaScript."""

import json
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "发现页内容数据库.xlsx"
OUTPUT = ROOT / "src" / "content-data.js"
SEP = "｜"
IMAGE_API = "https://copilot-cn.bytedance.net/api/ide/v1/text_to_image"


def split(value):
    return [part.strip() for part in str(value or "").split(SEP) if part.strip()]


def rows(ws):
    headers = [cell.value for cell in ws[1]]
    return [
        dict(zip(headers, [cell.value for cell in row]))
        for row in ws.iter_rows(min_row=2)
        if any(cell.value not in (None, "") for cell in row)
    ]


def image_url(purpose, name, image_size="landscape_16_9"):
    prompts = {
        "模板": (
            f"Website template cover for {name}: a realistic close-up of an enterprise project "
            "management workspace, structured workflow board, task cards, status columns and timeline, "
            "clean Chinese SaaS interface, light neutral canvas, blue accents, no readable text"
        ),
        "解决方案": (
            f"Website solution cover for {name}: a wide enterprise operating model showing connected "
            "strategy, project portfolio, milestones, delivery workflow and management dashboard, "
            "architectural overview blended with realistic software UI, navy and cyan accents, no readable text"
        ),
        "轻应用": (
            f"Website lightweight business app cover for {name}: a polished ready-to-use operations "
            "cockpit with KPI cards, charts, project risks, approvals and action list, realistic Chinese "
            "SaaS product UI, light theme with green accents, no readable text"
        ),
        "插件": (
            f"Website plugin cover for {name}: a clean enterprise software integration card showing "
            "connected systems, structured data flow and project task interface, realistic SaaS UI, "
            "light theme, no readable text"
        ),
    }
    prompt = prompts.get(purpose, prompts["模板"])
    return f"{IMAGE_API}?prompt={quote(prompt)}&image_size={image_size}"


def product_kind(value):
    return "AI 应用" if value == "AI应用" else value


def delivery_fields(relation, kind):
    result = {
        "implementationSummary": relation.get("使用方法概述") or "",
        "steps": split(relation.get("操作步骤")),
        "deliveryType": relation.get("交付类型") or "",
        "prompt": relation.get("参考提示词") or "",
        "prdContent": relation.get("PRD正文") or "",
        "prdUrl": relation.get("PRD链接") or "",
        "guideUrl": relation.get("搭建说明链接") or "",
        "ctaLabel": relation.get("按钮文案") or "",
        "configurationNotes": relation.get("配置注意事项") or "",
        "expectedOutput": relation.get("预期产出") or "",
    }
    target = relation.get("目标链接") or ""
    if kind == "模板":
        result["templateUrl"] = target
    elif kind == "解决方案":
        result["solutionUrl"] = target
    else:
        result["targetUrl"] = target
    return result


def compile_data():
    wb = load_workbook(SOURCE, data_only=True)
    scenario_rows = rows(wb["场景表"])
    product_rows = rows(wb["产品表"])
    relation_rows = rows(wb["场景产品关系表"])

    products_by_id = {row["产品ID"]: row for row in product_rows}
    relations_by_scenario = {}
    scenarios_by_product = {}
    for relation in relation_rows:
        relations_by_scenario.setdefault(relation["场景ID"], []).append(relation)
        scenarios_by_product.setdefault(relation["产品ID"], []).append(relation["场景ID"])
    scenarios_by_id = {row["场景ID"]: row for row in scenario_rows}

    scenarios = []
    for row in sorted(scenario_rows, key=lambda item: item.get("推荐排序") or 999):
        relations = relations_by_scenario.get(row["场景ID"], [])
        primary = next((item for item in relations if item.get("是否主要实现") == "是"), relations[0] if relations else {})
        product = products_by_id.get(primary.get("产品ID"), {})
        kind = product_kind(product.get("产品类型") or "AI应用")
        cover = row.get("封面图片") or product.get("封面图片") or ""
        if not cover and kind in {"插件", "模板", "解决方案", "轻应用"}:
            cover = image_url(kind if kind in {"插件", "模板", "解决方案", "轻应用"} else "轻应用", row["场景名称"])
        item = {
            "id": row["场景ID"],
            "title": row["场景名称"],
            "kind": kind,
            "workGoal": row.get("工作目标") or "",
            "subtype": product.get("产品子类型") or "",
            "featureName": product.get("产品名称") or "",
            "tag": kind,
            "industry": split(row.get("适用行业"))[0] if split(row.get("适用行业")) else "",
            "role": split(row.get("适用角色"))[0] if split(row.get("适用角色")) else "",
            "industries": split(row.get("适用行业")),
            "roles": split(row.get("适用角色")),
            "featured": row.get("是否推荐") == "是",
            "desc": row.get("卡片短描述") or "",
            "problem": row.get("问题描述") or "",
            "scenario": row.get("场景说明") or "",
            "capabilities": split(row.get("预期价值")),
            "tags": split(row.get("场景标签")),
            "image": cover,
            "customer": row.get("客户案例") or "",
            "metrics": row.get("量化结果") or "",
            "meta": " · ".join(split(row.get("场景标签"))[:2]),
            "contentStatus": row.get("内容状态") or "",
            "sourceType": "scenario",
        }
        item.update(delivery_fields(primary, kind))
        scenarios.append(item)

    products = []
    for row in sorted(product_rows, key=lambda item: (item.get("产品类型") or "", item.get("目录排序") or 999)):
        if row.get("目录是否展示") != "是":
            continue
        kind = product_kind(row["产品类型"])
        related_ids = scenarios_by_product.get(row["产品ID"], [])
        related = [scenarios_by_id[sid] for sid in related_ids if sid in scenarios_by_id]
        example = related[0] if related else None
        cover = row.get("封面图片") or ""
        if not cover and kind in {"模板", "解决方案", "插件"}:
            cover = image_url(kind, row["产品名称"])
        item = {
            "id": row["产品ID"],
            "title": row["产品名称"],
            "featureName": row["产品名称"],
            "kind": kind,
            "subtype": row.get("产品子类型") or "",
            "tag": kind,
            "industry": split(row.get("适用行业"))[0] if split(row.get("适用行业")) else "",
            "role": split(row.get("适用角色"))[0] if split(row.get("适用角色")) else "",
            "industries": split(row.get("适用行业")),
            "roles": split(row.get("适用角色")),
            "desc": row.get("卡片短描述") or "",
            "detailDescription": row.get("详情介绍") or "",
            "capabilities": split(row.get("核心能力")),
            "image": cover,
            "icon": row.get("图标") or "",
            "provider": row.get("提供方") or "",
            "updatedAt": row.get("最近更新时间") or "",
            "useUrl": row.get("直接使用链接") or "",
            "guideUrl": row.get("配置说明链接") or "",
            "productDocUrl": row.get("产品文档链接") or "",
            "sourceUrl": row.get("来源URL") or "",
            "exampleScenario": example.get("场景名称") if example else "",
            "exampleScenarioDesc": example.get("场景说明") if example else "",
            "relatedScenarios": [item.get("场景名称") for item in related],
            "contentStatus": row.get("内容状态") or "",
            "sourceType": "product",
        }
        if kind == "插件":
            item["prdUrl"] = row.get("产品文档链接") or ""
        elif kind == "模板":
            item["templateUrl"] = row.get("直接使用链接") or ""
        elif kind == "解决方案":
            item["solutionUrl"] = row.get("直接使用链接") or ""
        products.append(item)
    return scenarios, products


def write_module(scenarios, products):
    content = (
        "// Generated from 发现页内容数据库.xlsx. Do not edit manually.\n"
        f"export const scenarios = {json.dumps(scenarios, ensure_ascii=False, indent=2)};\n\n"
        f"export const products = {json.dumps(products, ensure_ascii=False, indent=2)};\n"
    )
    OUTPUT.write_text(content, encoding="utf-8")


def verify(scenarios, products):
    assert len(scenarios) >= 37
    assert len(products) >= 63
    assert all(item["featured"] for item in scenarios)
    assert all(item["kind"] != "轻应用" for item in products)
    assert sum(item["kind"] == "AI 应用" for item in products) >= 33
    assert sum(item["kind"] == "插件" for item in products) == 12
    assert sum(item["kind"] == "模板" for item in products) == 12
    assert sum(item["kind"] == "解决方案" for item in products) == 6
    assert sum(item["kind"] == "AI 应用" for item in scenarios) >= 31


if __name__ == "__main__":
    compiled_scenarios, compiled_products = compile_data()
    verify(compiled_scenarios, compiled_products)
    write_module(compiled_scenarios, compiled_products)
    print({
        "scenarios": len(compiled_scenarios),
        "products": len(compiled_products),
        "output": str(OUTPUT),
    })
